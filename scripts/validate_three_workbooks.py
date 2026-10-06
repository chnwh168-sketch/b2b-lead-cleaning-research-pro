#!/usr/bin/env python3

import argparse
import csv
import json
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


def norm(value):
    return unicodedata.normalize("NFKC", str(value or "")).strip()


def read_rows(path, id_column, email_column):
    workbook = load_workbook(path, read_only=True, data_only=False)
    worksheet = workbook.worksheets[0]
    header_row = None
    headers = []
    for row_number, values in enumerate(worksheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        candidate = [norm(value) for value in values]
        if id_column in candidate:
            header_row = row_number
            headers = candidate
            break
    if header_row is None:
        raise ValueError(f"{path}: cannot find header {id_column!r} in first 30 rows")

    id_index = headers.index(id_column)
    email_index = headers.index(email_column) if email_column in headers else None
    if email_index is None:
        raise ValueError(f"{path}: missing Email column")
    eligibility_index = headers.index("CRM Eligibility") if "CRM Eligibility" in headers else None
    if eligibility_index is None:
        raise ValueError(f"{path}: missing CRM Eligibility column")
    secondary_index = headers.index("Secondary_Email") if "Secondary_Email" in headers else None
    rows = []
    for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        company_id = norm(values[id_index] if id_index < len(values) else "")
        if not company_id:
            if any(norm(v) for v in values):
                raise ValueError(f"{path}: populated row has blank Company_ID")
            continue
        email = norm(values[email_index] if email_index is not None and email_index < len(values) else "").lower()
        rows.append({"company_id": company_id, "email": email,
                     "secondary": norm(values[secondary_index] if secondary_index is not None else "").lower(),
                     "eligibility": norm(values[eligibility_index])})
    return rows


def read_good_emails(path):
    if not path:
        return None
    with open(path, newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ValueError(f"{path}: empty CSV")
        field = next((name for name in reader.fieldnames if norm(name).lower() == "email"), None)
        if not field:
            raise ValueError(f"{path}: cannot find email column")
        return {norm(row.get(field)).lower() for row in reader if norm(row.get(field))}


def duplicates(values):
    seen = set()
    found = set()
    for value in values:
        if value in seen:
            found.add(value)
        seen.add(value)
    return sorted(found)


def main():
    parser = argparse.ArgumentParser(description="Validate the three mutually exclusive Zoho delivery workbooks.")
    parser.add_argument("--with-email", required=True, type=Path)
    parser.add_argument("--without-email", required=True, type=Path)
    parser.add_argument("--not-candidate", required=True, type=Path)
    parser.add_argument("--id-column", default="Company_ID")
    parser.add_argument("--email-column", default="Email")
    parser.add_argument("--expected-total", type=int)
    parser.add_argument("--good-only-csv", type=Path)
    parser.add_argument("--frozen-ids", required=True, type=Path, help="JSON array of frozen Company_ID values")
    args = parser.parse_args()

    groups = {
        "with_email": read_rows(args.with_email, args.id_column, args.email_column),
        "without_email": read_rows(args.without_email, args.id_column, args.email_column),
        "not_candidate": read_rows(args.not_candidate, args.id_column, args.email_column),
    }
    ids = {name: [row["company_id"] for row in rows] for name, rows in groups.items()}
    id_sets = {name: set(values) for name, values in ids.items()}
    good_emails = read_good_emails(args.good_only_csv)
    errors = []

    for name, values in ids.items():
        dupes = duplicates(values)
        if dupes:
            errors.append(f"{name}: duplicate Company_ID values: {dupes[:10]}")

    names = list(id_sets)
    for index, left in enumerate(names):
        for right in names[index + 1 :]:
            overlap = sorted(id_sets[left] & id_sets[right])
            if overlap:
                errors.append(f"{left}/{right}: overlapping Company_ID values: {overlap[:10]}")

    missing_email = [row["company_id"] for row in groups["with_email"] if not row["email"]]
    if missing_email:
        errors.append(f"with_email: blank Email for {missing_email[:10]}")

    unexpected_email = [row["company_id"] for row in groups["without_email"] if row["email"]]
    if unexpected_email:
        errors.append(f"without_email: Email must be blank for {unexpected_email[:10]}")

    for name in ("without_email", "not_candidate"):
        if any(row["email"] or row["secondary"] for row in groups[name]):
            errors.append(f"{name}: both email fields must be blank")
    for name in ("with_email", "without_email"):
        if any(row["eligibility"] != "Include" for row in groups[name]):
            errors.append(f"{name}: CRM Eligibility must be Include")
    if any(row["eligibility"] not in {"Exclude", "Manual Review"} for row in groups["not_candidate"]):
        errors.append("not_candidate: eligibility must be Exclude or Manual Review")
    frozen = json.loads(args.frozen_ids.read_text(encoding="utf-8"))
    if not isinstance(frozen, list) or any(not isinstance(x, str) or not norm(x) for x in frozen):
        raise ValueError("frozen IDs must be a nonblank string array")
    frozen = [norm(x) for x in frozen]
    if len(frozen) != len(set(frozen)):
        errors.append("duplicate frozen IDs")
    union = set().union(*id_sets.values())
    if union != set(frozen):
        errors.append(f"frozen set mismatch: missing={sorted(set(frozen)-union)[:10]}, extra={sorted(union-set(frozen))[:10]}")
    all_emails = [v for row in groups["with_email"] for v in (row["email"], row["secondary"]) if v]
    if duplicates(all_emails):
        errors.append("shared email requires explicit allocation policy; duplicate email found")
    if good_emails is None and all_emails:
        errors.append("Good-only CSV required for populated verified email fields")
    if good_emails is not None:
        not_good = sorted({v for v in all_emails if v not in good_emails})
        if not_good:
            errors.append(f"with_email: addresses absent from Good-only CSV: {not_good[:10]}")

    total = sum(len(values) for values in ids.values())
    if args.expected_total is not None and total != args.expected_total:
        errors.append(f"partition total {total} != expected {args.expected_total}")

    report = {
        "passed": not errors,
        "counts": {name: len(rows) for name, rows in groups.items()},
        "partitionTotal": total,
        "uniqueCompanyIds": len(set().union(*id_sets.values())),
        "errors": errors,
    }
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not errors else 1


if __name__ == "__main__":
    sys.exit(main())
